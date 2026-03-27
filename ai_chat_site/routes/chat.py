from __future__ import annotations

from datetime import datetime, timedelta, timezone
from zoneinfo import ZoneInfo
import os
import io
import mimetypes
from pathlib import Path

from flask import Blueprint, current_app, jsonify, render_template, request
from flask_login import login_required, current_user
from werkzeug.utils import secure_filename

from .. import limiter
from ..db import get_db
from ..gemini_service import generate_reply
from ..memory_service import recall, remember_message


bp = Blueprint("chat", __name__)

def _allowed_exts() -> set[str]:
    return {str(x).lower().strip().lstrip(".") for x in (current_app.config.get("UPLOAD_ALLOWED_EXT") or []) if str(x).strip()}

def _max_upload_bytes() -> int:
    return int(current_app.config.get("MAX_UPLOAD_BYTES") or 0)

def _upload_dir() -> str:
    return str(current_app.config.get("UPLOAD_DIR") or "/data/uploads")

def _is_image_mime(mime: str | None) -> bool:
    return bool(mime) and str(mime).lower().startswith("image/")

def _extract_text_from_bytes(data: bytes, filename: str, mime: str | None) -> str:
    ext = os.path.splitext(filename or "")[1].lower().lstrip(".")
    if mime and str(mime).lower().startswith("text/"):
        return data.decode("utf-8", errors="ignore")
    if ext in {"txt", "md", "csv", "json", "log"}:
        return data.decode("utf-8", errors="ignore")
    if ext == "pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            pages = []
            for p in reader.pages:
                pages.append(p.extract_text() or "")
            return "\n".join([t for t in pages if t.strip()])
        except Exception:
            return ""
    if ext == "docx":
        try:
            from docx import Document

            doc = Document(io.BytesIO(data))
            return "\n".join([p.text for p in doc.paragraphs if p.text and p.text.strip()])
        except Exception:
            return ""
    return ""

def _now_boundaries_utc() -> tuple[str, str]:
    tz_name = (current_app.config.get("TIMEZONE") or None) or (current_app.config.get("TZ") or None) or (os.getenv("TZ") or None) or "Asia/Shanghai"
    try:
        tz = ZoneInfo(str(tz_name))
    except Exception:
        tz = timezone.utc
    now_local = datetime.now(tz=tz)
    week_start_local = (now_local - timedelta(days=now_local.weekday())).replace(hour=0, minute=0, second=0, microsecond=0)
    month_start_local = now_local.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    week_start_utc = week_start_local.astimezone(timezone.utc)
    month_start_utc = month_start_local.astimezone(timezone.utc)
    return week_start_utc.strftime("%Y-%m-%d %H:%M:%S"), month_start_utc.strftime("%Y-%m-%d %H:%M:%S")


def _sum_tokens(user_id: int, *, conversation_id: int | None = None, since_utc: str | None = None) -> int:
    db = get_db()
    sql = "SELECT COALESCE(SUM(COALESCE(total_tokens,0)),0) AS s FROM chat_messages WHERE user_id=?"
    params: list = [int(user_id)]
    if conversation_id is not None:
        sql += " AND conversation_id=?"
        params.append(int(conversation_id))
    if since_utc is not None:
        sql += " AND created_at>=?"
        params.append(str(since_utc))
    row = db.execute(sql, tuple(params)).fetchone()
    return int(row["s"] or 0) if row else 0


def _get_or_create_default_conversation(user_id: int) -> int:
    db = get_db()
    row = db.execute(
        "SELECT id FROM conversations WHERE user_id=? ORDER BY updated_at DESC, id DESC LIMIT 1",
        (user_id,),
    ).fetchone()
    if row:
        return int(row["id"])
    cur = db.execute("INSERT INTO conversations(user_id, title) VALUES(?,?)", (user_id, "新对话"))
    db.commit()
    return int(cur.lastrowid)


def _require_conversation(user_id: int, conversation_id: int | None) -> int:
    if not conversation_id:
        return _get_or_create_default_conversation(user_id)
    db = get_db()
    row = db.execute(
        "SELECT id FROM conversations WHERE id=? AND user_id=?",
        (int(conversation_id), user_id),
    ).fetchone()
    if row:
        return int(row["id"])
    return _get_or_create_default_conversation(user_id)


def _fetch_history(user_id: int, conversation_id: int, limit: int = 20) -> list[dict]:
    db = get_db()
    rows = db.execute(
        "SELECT role, content FROM chat_messages WHERE user_id=? AND conversation_id=? ORDER BY id DESC LIMIT ?",
        (user_id, conversation_id, limit),
    ).fetchall()
    # Reverse to chronological.
    rows = list(reversed(rows))
    history = []
    for r in rows:
        role = "user" if r["role"] == "user" else "model"
        history.append({"role": role, "parts": [str(r["content"])]})
    return history


@bp.get("/chat")
@login_required
def index():
    db = get_db()
    conv_id = _get_or_create_default_conversation(int(current_user.id))
    conversations = db.execute(
        "SELECT id, title, updated_at FROM conversations WHERE user_id=? ORDER BY updated_at DESC, id DESC LIMIT 100",
        (int(current_user.id),),
    ).fetchall()
    allowed_models = current_app.config.get("GEMINI_ALLOWED_MODELS") or [current_app.config.get("GEMINI_MODEL")]
    return render_template(
        "chat/index.html",
        model_name=current_app.config.get("GEMINI_MODEL"),
        allowed_models=allowed_models,
        conversations=conversations,
        current_conversation_id=conv_id,
        memory_enabled_default=bool(current_app.config.get("MEMORY_ENABLED_DEFAULT", True)),
        upload_max_files=int(current_app.config.get("UPLOAD_MAX_FILES") or 5),
    )


@bp.get("/api/conversations")
@login_required
def api_conversations():
    db = get_db()
    rows = db.execute(
        """
        SELECT
          c.id,
          c.title,
          c.updated_at,
          COALESCE((
            SELECT SUM(COALESCE(m.total_tokens, 0))
            FROM chat_messages m
            WHERE m.user_id=c.user_id AND m.conversation_id=c.id
          ), 0) AS total_tokens
        FROM conversations c
        WHERE c.user_id=?
        ORDER BY c.updated_at DESC, c.id DESC
        LIMIT 200
        """,
        (int(current_user.id),),
    ).fetchall()
    return jsonify({"conversations": [dict(r) for r in rows]})

@bp.get("/api/stats")
@login_required
def api_stats():
    conv_id = request.args.get("conversation_id", type=int)
    conv_id = _require_conversation(int(current_user.id), conv_id)
    week_start_utc, month_start_utc = _now_boundaries_utc()
    return jsonify(
        {
            "conversation_id": conv_id,
            "current_chat_tokens": _sum_tokens(int(current_user.id), conversation_id=conv_id),
            "week_tokens": _sum_tokens(int(current_user.id), since_utc=week_start_utc),
            "month_tokens": _sum_tokens(int(current_user.id), since_utc=month_start_utc),
            "total_tokens": _sum_tokens(int(current_user.id)),
        }
    )

@bp.post("/api/upload")
@login_required
@limiter.limit("20 per minute")
def api_upload():
    file = request.files.get("file")
    if not file or not file.filename:
        return jsonify({"error": "未找到文件"}), 400

    filename = secure_filename(file.filename)
    if not filename:
        return jsonify({"error": "文件名无效"}), 400

    ext = os.path.splitext(filename)[1].lower().lstrip(".")
    allowed = _allowed_exts()
    if allowed and ext not in allowed:
        return jsonify({"error": "不支持的文件类型"}), 400

    data = file.read()
    if not data:
        return jsonify({"error": "文件为空"}), 400

    max_bytes = _max_upload_bytes()
    if max_bytes and len(data) > max_bytes:
        return jsonify({"error": "文件过大"}), 413

    mime = file.mimetype or mimetypes.guess_type(filename)[0] or "application/octet-stream"
    is_image = _is_image_mime(mime)
    extracted_text = "" if is_image else _extract_text_from_bytes(data, filename, mime)
    if extracted_text and len(extracted_text) > 200000:
        extracted_text = extracted_text[:200000]

    user_dir = Path(_upload_dir()) / str(int(current_user.id))
    user_dir.mkdir(parents=True, exist_ok=True)
    safe_name = f"{int(datetime.utcnow().timestamp())}_{filename}"
    path = user_dir / safe_name
    path.write_bytes(data)

    db = get_db()
    cur = db.execute(
        """
        INSERT INTO uploaded_files(user_id, original_name, storage_path, mime_type, size_bytes, extracted_text)
        VALUES(?,?,?,?,?,?)
        """,
        (int(current_user.id), filename, str(path), str(mime), int(len(data)), extracted_text),
    )
    db.commit()
    file_id = int(cur.lastrowid)
    return jsonify(
        {
            "id": file_id,
            "name": filename,
            "size": int(len(data)),
            "mime": str(mime),
            "is_image": bool(is_image),
        }
    )

@bp.delete("/api/upload/<int:file_id>")
@login_required
@limiter.limit("30 per minute")
def api_upload_delete(file_id: int):
    db = get_db()
    row = db.execute(
        "SELECT storage_path FROM uploaded_files WHERE id=? AND user_id=?",
        (int(file_id), int(current_user.id)),
    ).fetchone()
    if not row:
        return jsonify({"error": "文件不存在"}), 404
    try:
        Path(row["storage_path"]).unlink(missing_ok=True)
    except Exception:
        pass
    db.execute("DELETE FROM uploaded_files WHERE id=? AND user_id=?", (int(file_id), int(current_user.id)))
    db.commit()
    return jsonify({"ok": True})


@bp.post("/api/conversations")
@login_required
@limiter.limit("10 per minute")
def api_conversation_create():
    data = request.get_json(silent=True) or {}
    title = (data.get("title") or "").strip() or "新对话"
    if len(title) > 80:
        title = title[:80]
    db = get_db()
    cur = db.execute("INSERT INTO conversations(user_id, title) VALUES(?,?)", (int(current_user.id), title))
    db.commit()
    return jsonify({"id": int(cur.lastrowid), "title": title})


@bp.patch("/api/conversations/<int:conversation_id>")
@login_required
@limiter.limit("30 per minute")
def api_conversation_rename(conversation_id: int):
    data = request.get_json(silent=True) or {}
    title = (data.get("title") or "").strip()
    if not title:
        return jsonify({"error": "标题不能为空"}), 400
    if len(title) > 80:
        title = title[:80]
    db = get_db()
    cur = db.execute(
        "UPDATE conversations SET title=?, updated_at=CURRENT_TIMESTAMP WHERE id=? AND user_id=?",
        (title, conversation_id, int(current_user.id)),
    )
    db.commit()
    if cur.rowcount != 1:
        return jsonify({"error": "会话不存在"}), 404
    return jsonify({"ok": True, "title": title})


@bp.delete("/api/conversations/<int:conversation_id>")
@login_required
@limiter.limit("10 per minute")
def api_conversation_delete(conversation_id: int):
    db = get_db()
    cur = db.execute("DELETE FROM conversations WHERE id=? AND user_id=?", (conversation_id, int(current_user.id)))
    db.commit()
    if cur.rowcount != 1:
        return jsonify({"error": "会话不存在"}), 404
    return jsonify({"ok": True})


@bp.get("/api/conversations/<int:conversation_id>/messages")
@login_required
def api_conversation_messages(conversation_id: int):
    db = get_db()
    row = db.execute(
        "SELECT id FROM conversations WHERE id=? AND user_id=?",
        (conversation_id, int(current_user.id)),
    ).fetchone()
    if not row:
        return jsonify({"error": "会话不存在"}), 404
    rows = db.execute(
        """
        SELECT
          role,
          content,
          created_at,
          model_name,
          prompt_tokens,
          completion_tokens,
          total_tokens
        FROM chat_messages
        WHERE user_id=? AND conversation_id=?
        ORDER BY id ASC
        LIMIT 500
        """,
        (int(current_user.id), conversation_id),
    ).fetchall()
    return jsonify({"messages": [dict(r) for r in rows]})


@bp.post("/api/chat")
@login_required
@limiter.limit("15 per minute")
def api_chat():
    data = request.get_json(silent=True) or {}
    msg = (data.get("message") or "").strip()
    file_ids = data.get("file_ids") or []
    if not isinstance(file_ids, list):
        file_ids = []
    file_ids = [int(x) for x in file_ids if str(x).isdigit()]
    file_ids = list(dict.fromkeys(file_ids))
    max_files = int(current_app.config.get("UPLOAD_MAX_FILES") or 5)
    if max_files and len(file_ids) > max_files:
        return jsonify({"error": f"最多上传 {max_files} 个文件"}), 400
    if not msg and not file_ids:
        return jsonify({"error": "请输入内容或上传文件"}), 400
    if len(msg) > 4000:
        return jsonify({"error": "内容过长（最大 4000 字符）"}), 400

    api_key = current_app.config.get("GEMINI_API_KEY")
    if not api_key:
        return jsonify({"error": "服务端未配置 GEMINI_API_KEY"}), 500

    conversation_id = _require_conversation(int(current_user.id), data.get("conversation_id"))

    model_name = (data.get("model") or "").strip() or current_app.config.get("GEMINI_MODEL", "gemini-2.5-flash")
    allowed_models = current_app.config.get("GEMINI_ALLOWED_MODELS") or []
    if allowed_models and model_name.lower() not in set(allowed_models):
        model_name = current_app.config.get("GEMINI_MODEL", "gemini-2.5-flash")

    db = get_db()
    attachments: list[dict] = []
    file_names: list[str] = []
    if file_ids:
        placeholders = ",".join(["?"] * len(file_ids))
        rows = db.execute(
            f"""
            SELECT id, original_name, storage_path, mime_type, size_bytes, extracted_text
            FROM uploaded_files
            WHERE user_id=? AND id IN ({placeholders})
            """,
            (int(current_user.id), *file_ids),
        ).fetchall()
        for r in rows:
            name = str(r["original_name"] or "")
            mime = str(r["mime_type"] or "")
            file_names.append(name or "附件")
            if _is_image_mime(mime):
                try:
                    data_bytes = Path(r["storage_path"]).read_bytes()
                    attachments.append({"type": "image", "mime_type": mime, "data": data_bytes})
                except Exception:
                    continue
            else:
                text = str(r["extracted_text"] or "")
                if text:
                    if len(text) > 50000:
                        text = text[:50000] + "…"
                    attachments.append({"type": "text", "text": f"文件：{name}\n{text}"})

    if not msg and attachments:
        msg = "请分析我上传的文件。"

    msg_for_db = msg
    if file_names:
        shown = "、".join(file_names[:6])
        msg_for_db = (msg_for_db or "（上传了附件）") + f"\n[附件] {shown}"

    cur_user = db.execute(
        "INSERT INTO chat_messages(user_id, conversation_id, role, content, model_name) VALUES(?,?,?,?,?)",
        (current_user.id, conversation_id, "user", msg_for_db, model_name),
    )
    db.commit()
    user_message_id = int(cur_user.lastrowid or 0)

    history = _fetch_history(int(current_user.id), conversation_id, limit=20)
    # For Gemini history, exclude the latest user msg we just stored; we pass it separately.
    if history and history[-1]["role"] == "user" and history[-1]["parts"] and history[-1]["parts"][0] == msg:
        history = history[:-1]

    memory_enabled = data.get("memory_enabled")
    if memory_enabled is None:
        memory_enabled = bool(current_app.config.get("MEMORY_ENABLED_DEFAULT", True))
    memory_enabled = bool(memory_enabled)

    memory_snippets: list[str] = []
    if memory_enabled:
        try:
            hits = recall(
                user_id=int(current_user.id),
                api_key=api_key,
                embed_model=str(current_app.config.get("MEMORY_EMBED_MODEL") or "text-embedding-004"),
                query=msg,
                top_k=int(current_app.config.get("MEMORY_TOP_K") or 5),
            )
            for h in hits:
                t = (h.content or "").strip()
                if not t:
                    continue
                if len(t) > 400:
                    t = t[:400] + "…"
                if t not in memory_snippets:
                    memory_snippets.append(t)
        except Exception:
            memory_snippets = []

    try:
        reply = generate_reply(
            api_key=api_key,
            model_name=model_name,
            user_message=msg,
            history=history,
            attachments=attachments,
            memory_snippets=memory_snippets,
        )
    except Exception:
        return jsonify({"error": "Gemini 调用失败，请稍后重试"}), 502

    reply_text = (reply.text or "").strip()
    if len(reply_text) > 20000:
        reply_text = reply_text[:20000] + "\n\n（已截断）"

    cur_model = db.execute(
        """
        INSERT INTO chat_messages(
          user_id, conversation_id, role, content, model_name, prompt_tokens, completion_tokens, total_tokens
        ) VALUES(?,?,?,?,?,?,?,?)
        """,
        (
            current_user.id,
            conversation_id,
            "model",
            reply_text,
            reply.model_name,
            reply.usage.prompt_tokens,
            reply.usage.completion_tokens,
            reply.usage.total_tokens,
        ),
    )
    db.execute("UPDATE conversations SET updated_at=CURRENT_TIMESTAMP WHERE id=? AND user_id=?", (conversation_id, int(current_user.id)))
    db.commit()
    model_message_id = int(cur_model.lastrowid or 0)

    # Best-effort memory indexing (failure should not block chat).
    if memory_enabled:
        try:
            remember_message(
                user_id=int(current_user.id),
                role="user",
                content=msg,
                api_key=api_key,
                embed_model=str(current_app.config.get("MEMORY_EMBED_MODEL") or "text-embedding-004"),
                max_items=int(current_app.config.get("MEMORY_MAX_ITEMS") or 2000),
                source_conversation_id=conversation_id,
                source_message_id=user_message_id or None,
            )
            remember_message(
                user_id=int(current_user.id),
                role="model",
                content=reply_text,
                api_key=api_key,
                embed_model=str(current_app.config.get("MEMORY_EMBED_MODEL") or "text-embedding-004"),
                max_items=int(current_app.config.get("MEMORY_MAX_ITEMS") or 2000),
                source_conversation_id=conversation_id,
                source_message_id=model_message_id or None,
            )
        except Exception:
            pass

    return jsonify(
        {
            "reply": reply_text,
            "conversation_id": conversation_id,
            "model": reply.model_name,
            "memory_used": bool(memory_snippets),
            "usage": {
                "prompt_tokens": reply.usage.prompt_tokens,
                "completion_tokens": reply.usage.completion_tokens,
                "total_tokens": reply.usage.total_tokens,
            },
        }
    )


@bp.post("/api/chat/clear")
@login_required
@limiter.limit("5 per minute")
def api_clear():
    data = request.get_json(silent=True) or {}
    conversation_id = _require_conversation(int(current_user.id), data.get("conversation_id"))
    db = get_db()
    db.execute(
        "DELETE FROM chat_messages WHERE user_id=? AND conversation_id=?",
        (current_user.id, conversation_id),
    )
    db.execute("UPDATE conversations SET updated_at=CURRENT_TIMESTAMP WHERE id=? AND user_id=?", (conversation_id, int(current_user.id)))
    db.commit()
    return jsonify({"ok": True})
